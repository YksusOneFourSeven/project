"""Все форматы экспорта используют одинаковые строки и одинаковые права доступа."""

from io import BytesIO
from decimal import Decimal
from fastapi import HTTPException
from app.db import rows, one
from app.service import allow, assignment_access


async def report(conn, user, kind, group_id=None, assignment_id=None, student_id=None):
    role = user["role"]
    if role == "student":
        if kind not in ("student", "ranking"):
            raise HTTPException(403, "Доступны только собственные результаты")
        if student_id and student_id != user["student_id"]:
            raise HTTPException(403, "Чужая ведомость недоступна")
        student_id = user["student_id"]
        s = await one(conn, "SELECT group_id FROM students WHERE id=%s", (student_id,))
        if group_id and group_id != s["group_id"]:
            raise HTTPException(403, "Чужая группа недоступна")
        group_id = s["group_id"]
    if role == "teacher":
        if kind != "sheet":
            raise HTTPException(
                403, "Преподавателю доступна ведомость своей дисциплины"
            )
        if not assignment_id:
            raise HTTPException(422, "Выберите назначение")
        await assignment_access(conn, user, assignment_id)
    filters = []
    params = []
    for field, value in [("group_id", group_id), ("assignment_id", assignment_id)]:
        if value is not None:
            filters.append("r." + field + "=%s")
            params.append(value)
    base = " FROM results r JOIN groups gr ON gr.id=r.group_id"
    where = " WHERE " + " AND ".join(filters) if filters else ""
    if kind == "sheet":
        if not assignment_id:
            raise HTTPException(422, "Выберите назначение группы и дисциплины")
        data = await rows(
            conn,
            'SELECT r.full_name AS "Студент",r.discipline AS "Дисциплина",r.semester AS "Семестр",r.final_grade AS "Итог",r.incomplete AS "Не завершено"'
            + base
            + where
            + " ORDER BY r.full_name",
            params,
        )
    elif kind == "student":
        if not student_id:
            raise HTTPException(422, "Выберите студента")
        filters.append("r.student_id=%s")
        params.append(student_id)
        data = await rows(
            conn,
            'SELECT r.full_name AS "Студент",r.discipline AS "Дисциплина",r.academic_year AS "Год",r.semester AS "Семестр",r.final_grade AS "Итог",r.control_form AS "Контроль",r.incomplete AS "Не завершено"'
            + base
            + " WHERE "
            + " AND ".join(filters)
            + " ORDER BY r.academic_year,r.semester,r.discipline",
            params,
        )
    elif kind == "ranking":
        if not group_id:
            raise HTTPException(422, "Выберите группу")
        # Сначала считаем рейтинг всей группы, затем возвращаем студенту только его строку.
        data = await rows(
            conn,
            "WITH scores AS (SELECT r.student_id,r.full_name,round(avg(r.final_grade) FILTER(WHERE r.control_form<>'зачет'),2) avg_grade"
            + base
            + where
            + ' GROUP BY r.student_id,r.full_name), ranked AS (SELECT student_id,full_name,avg_grade,dense_rank() OVER(ORDER BY avg_grade DESC NULLS LAST) rank FROM scores) SELECT student_id AS "ID",full_name AS "Студент",avg_grade AS "Средний балл",CASE WHEN avg_grade IS NOT NULL THEN rank END AS "Место" FROM ranked'
            + (" WHERE student_id=%s" if role == "student" else "")
            + " ORDER BY rank,student_id",
            params + ([student_id] if role == "student" else []),
        )
    elif kind == "debtors":
        allow(user, "admin", "dean")
        filters.append(
            "((r.control_form<>'зачет' AND r.final_grade=2) OR (r.control_form='зачет' AND r.final_grade=0))"
        )
        data = await rows(
            conn,
            'SELECT gr.name AS "Группа",r.full_name AS "Студент",r.discipline AS "Дисциплина",r.final_grade AS "Итог"'
            + base
            + " WHERE "
            + " AND ".join(filters)
            + " ORDER BY gr.name,r.full_name",
            params,
        )
    elif kind == "statistics":
        allow(user, "admin", "dean")
        # GROUPING SETS возвращает независимые срезы, а не усреднение средних групп.
        data = await rows(
            conn,
            "SELECT CASE WHEN grouping(gr.name)=0 THEN 'группа' WHEN grouping(gr.course)=0 THEN 'курс' ELSE 'специальность' END AS \"Срез\",coalesce(gr.name,gr.course::text,gr.specialty) AS \"Значение\",count(DISTINCT r.student_id) AS \"Студентов\",count(r.final_grade) AS \"Итогов\",round(avg(r.final_grade) FILTER(WHERE r.control_form<>'зачет'),2) AS \"Средний балл\",round(100.0*count(*) FILTER(WHERE (r.control_form='зачет' AND r.final_grade=1) OR (r.control_form<>'зачет' AND r.final_grade>=3))/nullif(count(r.final_grade),0),2) AS \"Успеваемость %%\""
            + base
            + where
            + " GROUP BY GROUPING SETS ((gr.name),(gr.course),(gr.specialty)) ORDER BY 1,2",
            params,
        )
    elif kind == "quality":
        allow(user, "admin", "dean")
        # Качество = студенты без пропусков, со всеми балльными итогами 4/5 и всеми зачетами.
        data = await rows(
            conn,
            "WITH per_student AS (SELECT r.student_id,gr.name,bool_and(NOT r.incomplete AND CASE WHEN r.control_form='зачет' THEN r.final_grade=1 ELSE r.final_grade>=4 END) good,count(*) FILTER(WHERE r.control_form<>'зачет') numeric_count"
            + base
            + where
            + ' GROUP BY r.student_id,gr.name) SELECT name AS "Группа",count(*) AS "Студентов",count(*) FILTER(WHERE good AND numeric_count>0) AS "На 4 и 5",round(100.0*count(*) FILTER(WHERE good AND numeric_count>0)/count(*),2) AS "Качество %%" FROM per_student GROUP BY name ORDER BY name',
            params,
        )
    else:
        raise HTTPException(404, "Отчет не найден")
    return data


TITLES = dict(
    sheet="Ведомость по группе и дисциплине",
    student="Сводная ведомость студента",
    ranking="Рейтинг студентов группы",
    debtors="Отчет по должникам",
    statistics="Статистика успеваемости",
    quality="Качество знаний",
)


def export_xlsx(data):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Отчет"
    if data:
        ws.append(list(data[0]))
        for row in data:
            ws.append([float(v) if isinstance(v, Decimal) else v for v in row.values()])
            # Текст пользователя не должен интерпретироваться Excel как формула.
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="E3EAF1")
        for i in range(1, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(i)].width = 28
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes = "A2"
    else:
        ws.append(["Нет данных"])
    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def export_pdf(data, title):
    from reportlab.platypus import (
        SimpleDocTemplate,
        Table,
        TableStyle,
        Paragraph,
        Spacer,
    )
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.pagesizes import A4, landscape
    from xml.sax.saxutils import escape

    pdfmetrics.registerFont(
        TTFont("DejaVu", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
    )
    styles = getSampleStyleSheet()
    styles["Normal"].fontName = "DejaVu"
    styles["Normal"].fontSize = 8
    styles["Normal"].leading = 11
    styles["Title"].fontName = "DejaVu"
    styles["Title"].fontSize = 16
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 12)]
    if data:

        def p(v):
            return Paragraph(
                escape(
                    "—"
                    if v is None
                    else "да" if v is True else "нет" if v is False else str(v)
                ),
                styles["Normal"],
            )

        table = Table(
            [[p(x) for x in data[0]]] + [[p(x) for x in r.values()] for r in data],
            colWidths=[770 / len(data[0])] * len(data[0]),
            repeatRows=1,
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), "#e3eaf1"),
                    ("GRID", (0, 0), (-1, -1), 0.4, "#aaaaaa"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ]
            )
        )
        story.append(table)
    else:
        story.append(Paragraph("Нет данных", styles["Normal"]))
    output = BytesIO()
    SimpleDocTemplate(
        output, pagesize=landscape(A4), leftMargin=30, rightMargin=30
    ).build(story)
    return output.getvalue()
